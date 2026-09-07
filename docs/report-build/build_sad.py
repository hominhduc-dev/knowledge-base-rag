from pathlib import Path
import re, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
OUT.mkdir(exist_ok=True)
FONT = 'C:/Windows/Fonts/arial.ttf'
BOLD = 'C:/Windows/Fonts/arialbd.ttf'

def diagram(name, nodes, edges, size=(1500,850)):
    im=Image.new('RGB',size,'white'); d=ImageDraw.Draw(im)
    f=ImageFont.truetype(FONT,25)
    for a,b,label in edges:
        x,y,w,h,_=nodes[a]; X,Y,W,H,_=nodes[b]
        if abs(X-x)>abs(Y-y):
            start=(x+w if X>x else x,y+h/2); end=(X if X>x else X+W,Y+H/2)
        else:
            start=(x+w/2,y+h if Y>y else y); end=(X+W/2,Y if Y>y else Y+H)
        d.line([start,end],fill='#486478',width=3)
        theta=math.atan2(end[1]-start[1],end[0]-start[0])
        d.polygon([end,(end[0]-14*math.cos(theta-.4),end[1]-14*math.sin(theta-.4)),(end[0]-14*math.cos(theta+.4),end[1]-14*math.sin(theta+.4))],fill='#486478')
        if label:
            mx=(start[0]+end[0])/2; my=(start[1]+end[1])/2
            box=d.textbbox((mx,my),label,font=f,anchor='mm'); d.rectangle((box[0]-4,box[1]-4,box[2]+4,box[3]+4),fill='white');d.text((mx,my),label,font=f,fill='#243f52',anchor='mm')
    for x,y,w,h,text in nodes.values():
        if text.endswith('?'):
            d.polygon([(x+w/2,y),(x+w,y+h/2),(x+w/2,y+h),(x,y+h/2)],fill='#eff5f8',outline='#285b76',width=3)
        else:
            d.rounded_rectangle((x,y,x+w,y+h),radius=12,fill='#eff5f8',outline='#285b76',width=3)
        lines=text.split('\n')
        for i,line in enumerate(lines): d.text((x+w/2,y+h/2+(i-(len(lines)-1)/2)*33),line,font=f,fill='#193a50',anchor='mm')
    im.save(OUT/(name+'.png'))

diagram('architecture',{
 'ui':(30,100,320,150,'Giao diện Next.js\nĐăng nhập • Hỏi đáp\nKho tài liệu • Quản trị'),
 'api':(570,100,350,150,'API Express\nXác thực và phân quyền\nChat • Truy hồi • Tài liệu'),
 'db':(1130,100,340,150,'PostgreSQL + pgvector\nPrisma và SQL\nDữ liệu • Vector • Job'),
 'ai':(570,460,350,130,'Gemini API\nEmbedding và sinh trả lời'),
 'worker':(1130,460,340,130,'Worker nhập liệu\nĐọc tệp • Chia đoạn\nTạo embedding'),
 'file':(30,460,320,130,'Kho tệp uploads\nPDF và DOCX')},
 [('ui','api','HTTPS / SSE'),('api','db','Truy vấn'),('api','ai','HTTPS'),('worker','db','Đọc / ghi'),('worker','ai','Embedding'),('api','file','Tệp')],(1500,660))
diagram('activity',{
 'a':(500,20,480,80,'Sinh viên gửi câu hỏi'), 'b':(500,155,480,90,'Xác thực và lọc phạm vi CNTT'),
 'c':(500,300,480,90,'Truy hồi các đoạn liên quan'), 'd':(500,450,480,90,'Có nguồn phù hợp?'),
 'e':(40,610,420,100,'Thông báo không đủ nguồn'), 'f':(1010,610,450,100,'Sinh câu trả lời từ nguồn'),
 'g':(1010,790,450,100,'Kiểm tra marker trích dẫn'), 'h':(500,980,480,100,'Lưu và trả kết quả cuối cùng')},
 [('a','b',''),('b','c','Hợp lệ'),('c','d',''),('d','e','Không'),('d','f','Có'),('f','g',''),('g','h',''),('e','h','')],(1500,1120))
diagram('erd',{
 'u':(30,40,380,100,'users'), 'm':(550,40,380,100,'department_members'), 'd':(1080,40,380,100,'departments'),
 'v':(30,250,380,100,'conversations'), 'doc':(1080,250,380,100,'documents'),
 'msg':(30,460,380,100,'messages'), 'chunk':(1080,460,380,100,'chunks'),
 'cite':(550,460,380,100,'message_citations'), 'emb':(1080,680,380,100,'chunk_embeddings'),
 'job':(550,250,380,100,'ingest_jobs')},
 [('u','m','1 : N'),('d','m','1 : N'),('u','v','1 : N'),('v','msg','1 : N'),('d','doc','0..1 : N'),('doc','chunk','1 : N'),('chunk','emb','1 : N'),('doc','job','1 : N'),('msg','cite','1 : N'),('chunk','cite','0..1 : N')],(1500,820))
diagram('classes',{
 'b':(40,60,390,160,'Boundary\nChatBox / useSseStream\nHiển thị nguồn và trả lời'),
 'c':(560,60,390,160,'Control\nChatController\nTuần tự hóa SSE'),
 's':(1070,60,390,160,'Business\nChatService.hoi()\nĐiều phối lượt hỏi đáp'),
 'r':(1070,430,390,160,'RetrievalService.search()\nSQL truy hồi và scope'),
 'e':(560,430,390,160,'Entity / Data Access\nPrisma + SQL\nConversation • Chunk'),
 'ai':(1070,760,390,160,'Gemini adapter\nEmbedding • Generate')},
 [('b','c','Gọi'),('c','s','Gọi'),('s','r','Truy hồi'),('r','e','Đọc'),('s','e','Lưu'),('r','ai','Embedding')],(1500,960))
diagram('screen',{
 'a':(30,150,300,100,'Đăng nhập'), 'b':(500,150,350,100,'Hỏi đáp / Hội thoại'),
 'c':(1110,30,350,100,'Kho tài liệu'), 'd':(1110,230,350,100,'Quản lý tài liệu'),
 'e':(1110,430,350,100,'Quản lý tài khoản')},
 [('a','b','Thành công'),('b','c','Ba vai'),('b','d','Giáo vụ / QTV'),('b','e','QTV')],(1500,580))

# Low-fidelity design, intentionally distinguished from screenshots of the app.
im=Image.new('RGB',(1500,760),'white');d=ImageDraw.Draw(im); f=ImageFont.truetype(FONT,26)
d.rectangle((20,20,1480,740),outline='#38566a',width=3)
d.rectangle((20,20,1480,95),fill='#eaf2f7'); d.text((45,43),'SỔ TAY SINH VIÊN CNTT',font=f,fill='#193a50')
d.line((340,95,340,740),fill='#38566a',width=2)
for i,t in enumerate(['Hỏi đáp','Kho tài liệu','Hội thoại của tôi','Quản lý nội dung *','Quản lý tài khoản **']): d.text((45,140+i*65),t,font=f,fill='#193a50')
d.rounded_rectangle((390,145,1410,230),radius=12,outline='#38566a',width=2);d.text((415,175),'Câu hỏi của sinh viên về quy chế học vụ',font=f,fill='#193a50')
d.rounded_rectangle((390,275,1410,475),radius=12,fill='#eff5f8',outline='#38566a',width=2)
for i,t in enumerate(['Câu trả lời có căn cứ từ kho tài liệu','Nguồn [1] • Tên văn bản • Trang / Điều','Bấm nguồn để xem đoạn trích và tài liệu gốc']):d.text((415,310+i*50),t,font=f,fill='#193a50')
d.rectangle((390,585,1270,665),outline='#38566a',width=2);d.text((415,610),'Nhập câu hỏi…',font=f,fill='#193a50');d.rectangle((1300,585,1410,665),fill='#d8e9f3');d.text((1325,610),'Gửi',font=f,fill='#193a50')
im.save(OUT/'wireframe.png')

doc=Document(); sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.2);sec.right_margin=Cm(1.8)
for name in ['Normal','Body Text','List Bullet','List Number']:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(11)
    st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.12
for name,size in [('Title',25),('Heading 1',16),('Heading 2',13),('Heading 3',11)]:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('Sổ Tay Sinh Viên CNTT  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
list_num=0
def p(t,style=None):
    global list_num
    if style=='List Number':
        list_num+=1
        return doc.add_paragraph(f'{list_num}. {t}')
    return doc.add_paragraph(t,style)
next_page=False
def h(t,level=1):
    global next_page,list_num
    list_num=0
    pp=doc.add_heading(t,level)
    if next_page:pp.paragraph_format.page_break_before=True;next_page=False
def page():
    global next_page
    next_page=True
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    for c,v in zip(t.rows[0].cells,headers):c.text=v
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
        if ri==0:
            trPr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            if widths:c.width=Cm(widths[ci])
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.space_before=Pt(2);pp.paragraph_format.line_spacing=1
                if ri==0 or (headers[0]=='Tên trường' and ri<len(t.rows)-1):pp.paragraph_format.keep_with_next=True
                for r in pp.runs:r.font.size=Pt(9);r.bold=(ri==0)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF1F5');c._tc.get_or_add_tcPr().append(shade)
def fig(path,caption,width=16.5):
    pp=doc.add_paragraph();pp.alignment=1;pp.paragraph_format.keep_with_next=True
    pp.add_run().add_picture(str(path),width=Cm(width))
    cp=p(caption,'Caption');cp.alignment=1

p('PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG','Title')
p('Sổ Tay Sinh Viên CNTT','Subtitle')
p('Học phần Lập trình mạng máy tính\nTrường Đại học Kiến trúc Đà Nẵng')
table(['Thông tin','Nội dung'],[['Phiên bản','1.0'],['Ngày thực hiện','05/09/2026'],['Đối tượng','Sinh viên ngành Công nghệ thông tin'],['Nhóm thực hiện','Nhóm đồ án học phần Lập trình mạng']],[4,13])
p('Tài liệu mô tả yêu cầu, kiến trúc, dữ liệu, giao diện và kế hoạch kiểm thử của hệ thống hỏi đáp học vụ dựa trên tài liệu. Hệ thống phục vụ phạm vi ngành CNTT, với ba vai trò Sinh Viên CNTT, Giáo vụ khoa CNTT và Quản Trị Viên. Câu trả lời cần có nguồn trích dẫn để người học kiểm chứng trong văn bản gốc.')
h('Nội dung tài liệu',2)
for t in ['1. Giới thiệu tổng quan hệ thống','2. Đặc tả yêu cầu hệ thống','3. Thiết kế kiến trúc và logic hệ thống','4. Thiết kế dữ liệu và lớp','5. Thiết kế giao diện','6. Kế hoạch triển khai và kiểm thử']:p(t)
page();h('1. GIỚI THIỆU TỔNG QUAN HỆ THỐNG')
h('1.1. Mục tiêu và lý do xây dựng dự án',2)
p('Thông tin học vụ thường nằm trong nhiều văn bản PDF và DOCX, khiến sinh viên mất thời gian tìm đúng điều khoản và xác định nguồn áp dụng. Giáo vụ phải giải đáp lặp lại các câu hỏi về quy chế, học phí, chương trình đào tạo và thủ tục học tập. Hệ thống tập trung tài liệu liên quan đến CNTT, hỗ trợ hỏi bằng tiếng Việt và dẫn lại vị trí nguồn.')
p('Mục tiêu là rút ngắn thao tác tra cứu, giúp sinh viên tự kiểm chứng câu trả lời và hỗ trợ giáo vụ cập nhật kho nội dung. Trong học phần Lập trình mạng, sản phẩm đồng thời minh họa giao tiếp client–server, REST qua HTTP, truyền kết quả bằng SSE, xác thực và xử lý lỗi kết nối.')
h('1.2. Phạm vi hệ thống',2)
table(['Thuộc phạm vi','Ngoài phạm vi'],[
 ['Hỏi đáp có nguồn, tìm kiếm tài liệu, xem hội thoại cá nhân.','Quản lý đào tạo toàn trường, đăng ký tín chỉ, điểm số, thanh toán.'],
 ['Tải lên, sửa thông tin, gỡ tài liệu và chạy lại xử lý lỗi.','Quản trị nhiều khoa; tự đăng ký hoặc SSO trong phiên bản hiện hành.'],
 ['Quản lý trạng thái và vai trò tài khoản trong phạm vi CNTT.','OCR tự động trong luồng nhập liệu thông thường; chưa coi PDF scan là đầu vào được bảo đảm.'],
 ['Đánh giá truy hồi, kiểm tra sức khỏe, thực hành TCP/HTTP trong LAN.','Cam kết quy mô phục vụ toàn trường hoặc tính sẵn sàng cấp sản phẩm thương mại.']],[8.5,8.5])
h('1.3. Đối tượng sử dụng',2)
table(['Actor','Vai kỹ thuật','Trách nhiệm'],[['Sinh Viên CNTT','USER','Hỏi đáp, đọc tài liệu, xem nguồn và hội thoại của mình.'],['Giáo vụ khoa CNTT','CONTENT_ADMIN','Có quyền tra cứu và quản lý nội dung tài liệu.'],['Quản Trị Viên','SYSTEM_ADMIN','Có quyền của giáo vụ và quản lý tài khoản, phân quyền.']],[4,4,9])
p('Cả ba vai chỉ truy cập tài liệu CNTT và quy định chung. Vai quản trị không bỏ qua giới hạn này. Gemini API là dịch vụ ngoài; người vận hành CLI là tác nhân kỹ thuật của công cụ đánh giá, không tạo thêm vai nghiệp vụ.')

page();h('2. ĐẶC TẢ YÊU CẦU HỆ THỐNG');h('2.1. Yêu cầu chức năng',2)
p('Danh mục giữ nguyên mã use case trong bộ sơ đồ hiện hành. UC17 không được sử dụng. Các chức năng API hoặc CLI được ghi rõ để phân biệt với màn hình đã có.')
rows=[]
readme=(ROOT/'docs/use-cases/README.md').read_text(encoding='utf8')
for line in readme.splitlines():
    if re.match(r'\| UC\d+',line):
        cols=[x.strip() for x in line.split('|')[1:-1]];rows.append(cols[:3])
table(['Mã','Chức năng','Giao tiếp và giới hạn'],rows,[1.3,5,10.7])
p('Không có API tự đăng ký hoặc POST /users; tài khoản được chuẩn bị bằng quy trình seed/bootstrap phục vụ triển khai. Đổi mật khẩu có API, không mặc định xem đây là màn hình đã hoàn thiện.')
h('Quy tắc nghiệp vụ',2)
for t in ['BR01. Tài khoản phải đang hoạt động và có tư cách CNTT để đăng nhập.','BR02. Mỗi request kiểm tra lại quyền từ CSDL; thay đổi vai hoặc khóa tài khoản tác động lên request tiếp theo.','BR03. Sinh viên chỉ đọc; giáo vụ và quản trị viên mới được ghi nội dung; chỉ quản trị viên thay đổi tài khoản.','BR04. Không tự khóa, tự hạ quyền hoặc tự gỡ tư cách của quản trị viên đang thao tác.','BR05. Hội thoại thuộc riêng người tạo. Truy cập nguồn và file gốc phải kiểm tra phạm vi.','BR06. Không có nguồn phù hợp hoặc không có trích dẫn hợp lệ thì trả thông báo từ chối thay vì khẳng định thiếu căn cứ.']:p(t)

page();h('2.2. Yêu cầu phi chức năng',2)
p('Các ngưỡng dưới đây là mục tiêu nghiệm thu đề xuất cho đồ án; cần đo trong môi trường demo trước khi công bố kết quả.')
table(['Nhóm','Yêu cầu và cách đánh giá'],[
 ['Hiệu năng','Mục tiêu p95 dưới 2 giây cho thao tác thông thường trên LAN, không tính sinh câu trả lời từ Gemini. Đo riêng độ trễ truy hồi, thời gian token đầu và tổng lượt hỏi đáp.'],
 ['Tải đồng thời','Bắt đầu kiểm thử với 10 client đồng thời trên bộ dữ liệu demo; ghi tỷ lệ lỗi và độ trễ. Không dùng mức 1.000 người dùng trong template làm cam kết.'],
 ['Bảo mật','Bcrypt cho mật khẩu; JWT cho API; kiểm tra quyền và scope tại backend. Dùng HTTPS ở môi trường công bố; không đưa khóa bí mật vào trình duyệt.'],
 ['Toàn vẹn','Không nhân đôi chunk khi có nhiều model embedding; lọc model hiện hành và scope trước khi xếp hạng. File trùng được nhận biết bằng hash.'],
 ['Khả năng phục hồi','Hiển thị lỗi mạng; kết thúc SSE có kiểm soát; theo dõi job FAILED và cho phép người có quyền chạy lại.'],
 ['Giao diện','Tiếng Việt; dùng được trên desktop và màn hình nhỏ; phân biệt trạng thái tải, lỗi, không có kết quả; nguồn có thể mở để kiểm chứng.'],
 ['Bảo trì','Phân tách giao diện, nghiệp vụ và truy cập dữ liệu; cấu hình môi trường tập trung; lưu migration và bộ kiểm thử trong repo.']],[3,14])
h('2.3. Biểu đồ Use Case tổng quát',2)
fig(OUT/'usecase.png','Hình 1. Use case tổng quát với ba actor nghiệp vụ',16.5)

page();h('2.4. Chi tiết đặc tả Use Case mẫu',2)
h('UC05 Hỏi đáp học vụ',3)
table(['Thuộc tính','Đặc tả'],[['Tác nhân chính','Sinh Viên CNTT; giáo vụ và quản trị viên cũng được tra cứu.'],['Điều kiện tiên quyết','Đã đăng nhập bằng tài khoản hoạt động có tư cách CNTT.'],['Kích hoạt','Người dùng nhập câu hỏi và chọn Gửi.'],['Hậu điều kiện','Câu hỏi, kết quả cuối và nguồn sử dụng được lưu trong hội thoại của người dùng khi lượt xử lý hoàn tất.']],[4,13])
h('Luồng sự kiện chính',3)
for t in ['Sinh viên nhập câu hỏi học vụ trên giao diện hỏi đáp.','Giao diện gửi POST /api/chat kèm token và mở luồng SSE.','Backend kiểm tra tài khoản, tư cách CNTT và quyền sở hữu hội thoại nếu tiếp tục cuộc trò chuyện.','Hệ thống truy hồi lai các đoạn thuộc tài liệu CNTT hoặc quy định chung, theo model embedding hiện hành.','Backend gửi sự kiện sources, sau đó dùng nguồn tìm được để yêu cầu Gemini sinh câu trả lời.','Giao diện nhận các sự kiện token; backend kiểm tra marker trích dẫn và xác định nội dung cuối.','Hệ thống lưu hội thoại, tin nhắn và trích dẫn. Sự kiện done mang nội dung cuối cùng, mã hội thoại và danh sách nguồn thực sự được dẫn.','Giao diện thay nội dung tạm bằng done.text; sinh viên mở nguồn để kiểm chứng.']:p(t,'List Number')
h('Luồng ngoại lệ',3)
for t in ['3a. Token không hợp lệ hoặc tài khoản bị khóa: trả lỗi xác thực; không tiếp tục truy hồi.','3b. Hội thoại không thuộc người gọi: không cung cấp nội dung hội thoại.','4a. Không có nguồn phù hợp: trả thông báo không đủ căn cứ.','5a. Gemini hoặc kết nối gặp lỗi: thông báo lỗi của lượt hỏi; không coi nội dung tạm là đáp án hoàn tất.','6a. Không có marker trích dẫn hợp lệ: thay kết quả bằng câu từ chối.','8a. Tài liệu gốc đã bị gỡ: lịch sử vẫn giữ bản sao trích dẫn, liên kết mở nguồn có thể không còn.']:p(t)

page();h('3. THIẾT KẾ KIẾN TRÚC VÀ LOGIC HỆ THỐNG');h('3.1. Kiến trúc tổng thể',2)
p('Hệ thống dùng mô hình client–server với ba lớp logic. Boundary gồm giao diện và HTTP client; Control gồm controller, service và kiểm soát quyền; Entity/Data Access gồm Prisma, truy vấn SQL và các thực thể PostgreSQL. Ba lớp logic không đồng nghĩa với ba máy chủ vật lý.')
fig(OUT/'architecture.png','Hình 2. Kiến trúc tổng thể và các kênh trao đổi')
table(['Lớp','Nội dung'],[['Presentation','Next.js/React hiển thị hỏi đáp, tài liệu, nguồn và tài khoản; nhận SSE qua client.'],['Business Logic','Express route/controller, service hỏi đáp, truy hồi và quản lý tài liệu; middleware kiểm quyền.'],['Data Access','Prisma cho nghiệp vụ thông thường; SQL cho truy hồi vector/full-text, PostgreSQL cho dữ liệu và hàng đợi.']],[4,13])
p('Worker xử lý tài liệu bất đồng bộ bằng hàng đợi ingest_jobs. Kho uploads lưu tệp; Gemini cung cấp embedding và sinh nội dung. Caddy làm reverse proxy khi triển khai Docker Compose. Bộ netlab riêng minh họa TCP framing và HTTP trong học phần.')
page();h('3.2. Biểu đồ hoạt động',2)
p('Luồng hỏi đáp chỉ đi đến sinh câu trả lời khi có nguồn. Lỗi xác thực kết thúc trước truy hồi; kết quả cuối luôn qua bước kiểm tra trích dẫn.')
fig(OUT/'activity.png','Hình 3. Luồng hoạt động hỏi đáp học vụ')
p('Với nhập liệu, tài liệu đi qua PENDING → PROCESSING → READY hoặc FAILED. Giáo vụ theo dõi trạng thái; thao tác chạy lại lỗi đưa tài liệu vào quy trình xử lý phù hợp thay vì yêu cầu sinh viên tự sửa dữ liệu.')
page();h('3.3. Biểu đồ tuần tự',2)
p('Mỗi use case có một sơ đồ riêng trong docs/use-cases/sequences. Hình 4 minh họa UC05 có tích hợp Gemini; các biểu tượng Actor, Boundary, Control và Entity dùng kiểu BCE đã thống nhất.')
seq=ROOT/'docs/use-cases/sequences/uc05_usecase_hoi_dap_hoc_vu/diagram.visual-check.2048x1320.light.png'
im=Image.open(seq);im.crop((40,85,2010,1080)).save(OUT/'sequence.png')
fig(OUT/'sequence.png','Hình 4. Sequence UC05 theo ba lớp và dịch vụ Gemini')
h('Dữ liệu trao đổi trong luồng SSE',3)
table(['Sự kiện','Nội dung và cách xử lý'],[['sources','Danh sách nguồn truy hồi; giao diện có thể hiển thị trước câu trả lời.'],['token','Phần văn bản tạm nhận được trong lượt sinh; chưa phải nội dung cuối.'],['done','messageId, conversationId, latencyMs, cited và text; giao diện thay câu trả lời tạm bằng text.'],['error','Mã lỗi và thông báo; kết thúc lượt thất bại có kiểm soát.']],[3,14])
p('Trình tự thông thường là sources → token (lặp nhiều lần) → done. Việc lưu nội dung và kiểm tra trích dẫn thuộc lớp nghiệp vụ; controller chịu trách nhiệm định dạng SSE. Entity không gửi kết quả trực tiếp tới actor.')

page();h('4. THIẾT KẾ DỮ LIỆU VÀ LỚP');h('4.1. Sơ đồ thực thể kết hợp',2)
fig(OUT/'erd.png','Hình 5. Các quan hệ chính của dữ liệu nghiệp vụ')
p('Quan hệ 1:N cho phép phía N có nhiều bản ghi; khóa ngoại có thể rỗng được ghi 0..1. Hình lược bớt đường users → documents, departments → chunks và documents → message_citations để dễ đọc. Các ràng buộc đầy đủ nằm trong từ điển dữ liệu bên dưới.')
p('Nhóm đánh giá gồm eval_sets, eval_questions, eval_gold_chunks, eval_runs và eval_results. Một bộ đánh giá có nhiều câu hỏi và nhiều lần chạy; mỗi câu hỏi liên kết các chunk kỳ vọng; kết quả nối một lần chạy với một câu hỏi. Nhóm này phục vụ công cụ CLI, không phải phân hệ quản trị web.')
h('Nguyên tắc dữ liệu',3)
for t in ['Vai trò nằm ở department_members.role, không nằm trực tiếp trong users. Ứng dụng chỉ sử dụng tư cách CNTT; cấu trúc departments được giữ để tương thích dữ liệu.','documents.department_id rỗng biểu thị tài liệu chung; khác rỗng phải khớp phạm vi CNTT. chunks lặp department_id và visibility để lọc trước xếp hạng; trigger đồng bộ từ tài liệu.','chunk_embeddings tách vector theo model; cặp chunk_id và model là duy nhất. Mọi truy hồi cần lọc model hiện hành.','message_citations lưu bản sao quote, heading_path và page; khi nguồn bị gỡ, khóa ngoại nguồn có thể thành NULL nhưng lịch sử trích dẫn vẫn còn.']:p(t)

h('4.2. Chi tiết cấu trúc các bảng cơ sở dữ liệu',2)
p('Từ điển sau dùng tên cột vật lý theo schema.prisma. PK là khóa chính; FK là khóa ngoại; UQ là duy nhất. Kiểu enum sử dụng đúng danh sách khai báo. Cột content_tsv được migration tạo tự động; dấu nullable trong Prisma phản ánh khai báo ORM, không thay thế việc kiểm tra DDL triển khai.')
schema=(ROOT/'server/prisma/schema.prisma').read_text(encoding='utf8')
models=re.findall(r'model (\w+) \{(.*?)\n\}',schema,re.S)
names={name:re.search(r'@@map\("([^"]+)"\)',body)[1] for name,body in models}
enum_names={'MemberRole':'member_role','DeptType':'dept_type','DocStatus':'doc_status','JobStatus':'job_status','MessageRole':'message_role'}
for name,body in models:
    fields=[];fk={}
    for line in body.splitlines():
        rel=re.search(r'@relation\([^\n]*fields: \[(\w+)\], references: \[(\w+)\]',line)
        if rel:
            target=line.split()[1].rstrip('?');fk[rel[1]]=names.get(target,target)+'.'+rel[2]
    for raw in body.splitlines():
        line=raw.strip()
        match=re.match(r'(\w+)\s+(Unsupported\("[^"]+"\)\??|\w+\??)(.*)',line)
        if not match or line.startswith('//'):continue
        field,typ,attrs=match.groups()
        if typ.rstrip('?') in names or '[]' in attrs:continue
        mapped=re.search(r'@map\("([^"]+)"\)',attrs);col=mapped[1] if mapped else field
        dbtype=re.search(r'@db\.(\w+(?:\([^)]*\))?)',attrs)
        dtype=dbtype[1] if dbtype else enum_names.get(typ.rstrip('?'),{'String':'text','Int':'integer','Boolean':'boolean','Float':'double precision','Json':'jsonb','DateTime':'timestamp'}.get(typ.rstrip('?'),typ.rstrip('?')))
        if typ.startswith('Unsupported'):dtype=re.search(r'"([^"]+)"',typ)[1]
        key='PK' if '@id' in attrs else ('FK' if field in fk else ('UQ' if '@unique' in attrs else '—'))
        note=[]
        if field in fk:note.append('→ '+fk[field])
        default=re.search(r'@default\(([^\n]+?)\)(?:\s|$)',attrs)
        if default:note.append('Mặc định '+default[1])
        if '@updatedAt' in attrs:note.append('Tự cập nhật')
        if col=='password_hash':note.append('Hash bcrypt')
        if col=='content_tsv':note.append('Sinh tự động từ content')
        fields.append([col,dtype,key,'Có' if typ.endswith('?') else 'Không','; '.join(note) or '—'])
    h('Bảng '+names[name],3)
    table(['Tên trường','Kiểu dữ liệu','Khóa','NULL','Ghi chú'],fields,[3.8,3.1,1.5,1.7,6.9])
    constraints=re.findall(r'@@(unique|index)\(([^\n]+)\)',body)
    if constraints:p('Ràng buộc / chỉ mục: '+ '; '.join(k+' '+v for k,v in constraints)+'.')
p('Enum: member_role = USER, CONTENT_ADMIN, SYSTEM_ADMIN; dept_type = FACULTY, OFFICE; doc_status = PENDING, PROCESSING, READY, FAILED; job_status = PENDING, PROCESSING, DONE, FAILED; message_role = USER, ASSISTANT.')
p('Migration còn quản lý chỉ mục HNSW cho embedding, GIN cho content_tsv, duy nhất lower(email) và trigger đồng bộ scope. Dùng db:check để kiểm chứng sau triển khai; không suy ra các cấu trúc này chỉ từ Prisma.')
h('4.3. Biểu đồ lớp',2)
p('Backend TypeScript tổ chức bằng module và hàm, không phải toàn bộ đều là class OOP. Hình dưới là mô hình trách nhiệm và phụ thuộc tương đương Boundary–Control–Entity; không khẳng định có lớp Repository riêng hoặc quan hệ kế thừa chưa tồn tại trong mã nguồn.')
fig(OUT/'classes.png','Hình 6. Trách nhiệm và phụ thuộc của các thành phần hỏi đáp')
table(['Thành phần','Trách nhiệm và dữ liệu trao đổi'],[['ChatController','Nhận AskInput, gọi service và ghi sự kiện SSE.'],['ChatService.hoi','AsyncGenerator sinh sources, token, done hoặc error; điều phối truy hồi, sinh và lưu kết quả.'],['RetrievalService.search','Trả Source[] với đoạn nguồn đã lọc phạm vi và model.'],['Prisma / SQL','Đọc ghi thực thể; thực hiện truy vấn PostgreSQL và bảo đảm ràng buộc.'],['Gemini adapter','Chuyển văn bản thành embedding hoặc sinh câu trả lời từ ngữ cảnh.']],[5,12])

page();h('5. THIẾT KẾ GIAO DIỆN');h('5.1. Sơ đồ di chuyển màn hình',2)
fig(OUT/'screen.png','Hình 7. Luồng màn hình theo vai trò')
p('Từ màn hình hỏi đáp, người dùng mở lịch sử của mình hoặc kho tài liệu. Khi chọn nguồn, ngăn trích dẫn hiển thị đoạn văn và đường dẫn tài liệu gốc. Menu nội dung chỉ xuất hiện cho giáo vụ và quản trị viên; menu tài khoản chỉ dành cho quản trị viên. Đăng xuất xóa phiên cục bộ và trở về đăng nhập.')
h('5.2. Bản vẽ phác thảo giao diện',2)
fig(OUT/'wireframe.png','Hình 8. Wireframe màn hình hỏi đáp, minh họa bố cục thiết kế')
p('* Giáo vụ và quản trị viên. ** Chỉ quản trị viên. Wireframe là bản phác thảo, không phải ảnh chụp giao diện đang chạy.')
table(['Màn hình','Thành phần và trạng thái'],[['Đăng nhập','Mã/email, mật khẩu, nút đăng nhập; lỗi xác thực, trạng thái đang gửi.'],['Hỏi đáp','Ô nhập, danh sách tin nhắn, nguồn, trạng thái stream; thay bằng kết quả cuối khi done.'],['Tài liệu','Danh sách, chi tiết và trạng thái xử lý; tải lên và thao tác nội dung theo quyền.'],['Tài khoản','Tìm người dùng, trạng thái active, vai trò; khóa thao tác tự hạ quyền/khóa chính mình.']],[4,13])

page();h('6. KẾ HOẠCH TRIỂN KHAI VÀ KIỂM THỬ');h('6.1. Môi trường và công nghệ sử dụng',2)
table(['Thành phần','Công nghệ trong repo'],[['Frontend','Next.js 16, React 19, TypeScript, Tailwind CSS.'],['Backend','Node.js, Express 5, TypeScript; Zod kiểm tra đầu vào, JWT và bcryptjs.'],['Dữ liệu','PostgreSQL 16 + pgvector; Prisma 6.16.3; SQL chuyên biệt cho truy hồi.'],['AI và nhập liệu','Gemini API; unpdf cho PDF có text và mammoth cho DOCX; worker theo hàng đợi PostgreSQL.'],['Hạ tầng','Docker Compose, Caddy, volume uploads và pgdata; có cấu hình PostgreSQL local hoặc cloud.'],['Thực hành mạng','TCP server/client, framing, HTTP server trong server/src/netlab; đo đạc LAN và quan sát REST/SSE.']],[4,13])
h('Trình tự triển khai đề xuất',3)
for t in ['Chuẩn bị cấu hình .env ở gốc, chọn DB demo riêng; không đưa khóa bí mật vào báo cáo hay source client.','Build dịch vụ, sinh Prisma client và chạy migration theo script db:deploy. Sao lưu trước khi thay đổi DB có dữ liệu.','Chuẩn bị tài khoản ba vai và kho tài liệu mẫu; chỉ dùng seed với DB demo, không dùng để nâng cấp dữ liệu thật.','Khởi động api, web, caddy và db nếu dùng profile local; kiểm tra health và kết nối Gemini bằng yêu cầu thực tế.','Đăng nhập bằng từng vai, thử hỏi đáp và quản lý nội dung, sau đó đo tải/độ trễ trong LAN. Ghi cấu hình và kết quả vào biên bản nghiệm thu.']:p(t,'List Number')
h('6.2. Kế hoạch kiểm thử',2)
p('Các trường hợp sau là kế hoạch kiểm chứng, không phải bảng kết quả chạy trong phiên lập tài liệu. Kiểm thử ghi dữ liệu phải dùng DB thử nghiệm riêng.')
table(['Mã','Mức','Kịch bản','Kết quả mong đợi'],[
 ['T01','Unit','Cắt đoạn, lọc marker trích dẫn.','Không tạo nguồn giả; giữ thông tin vị trí.'],
 ['T02','Integration','Cùng chunk có vector của hai model.','Không nhân đôi nguồn; chỉ dùng model hiện hành.'],
 ['T03','Integration','Truy hồi CNTT, chung và khoa khác.','Chỉ CNTT/chung, kể cả SYSTEM_ADMIN.'],
 ['T04','System','Đăng nhập đúng/sai mật khẩu; bị khóa; không có tư cách CNTT.','Chỉ tài khoản hợp lệ được vào.'],
 ['T05','Integration','Giữ JWT, thay đổi vai trong DB.','Request kế tiếp dùng quyền mới.'],
 ['T06','System','Sinh viên ghi tài liệu; giáo vụ quản lý tài khoản.','Backend từ chối dù gọi trực tiếp API.'],
 ['T07','System','Đọc hội thoại hoặc file ngoài quyền.','Không rò dữ liệu.'],
 ['T08','System','Tải tệp, theo dõi job; nhập trùng; retry lỗi.','Trạng thái nhất quán; không lặp tài liệu trùng.'],
 ['T09','System','Hỏi có nguồn / không nguồn / marker sai.','Trả lời có dẫn nguồn hoặc từ chối.'],
 ['T10','System','Mất mạng hoặc lỗi Gemini giữa SSE.','Thông báo lỗi; không coi bản tạm là kết quả cuối.'],
 ['T11','System','Tự khóa hoặc tự hạ vai quản trị.','Thao tác bị chặn.'],
 ['T12','Performance','10 client LAN, tác vụ đọc và hỏi đáp.','Ghi p50/p95, tỷ lệ lỗi, token đầu và tổng độ trễ.'],
 ['T13','Evaluation','Chạy bộ câu hỏi cntt-v1.','Báo cáo recall@k, MRR và độ trễ theo cấu hình.'],
 ['T14','Network','TCP chia/gộp frame; nhiều client; quan sát HTTP/SSE.','Đọc đúng thông điệp, xử lý ngắt kết nối.']],[1,2.4,7.3,6.3])
h('Điều kiện nghiệm thu',3)
p('Không có lỗi rò phạm vi, vượt quyền hoặc lẫn hội thoại. Các thao tác cốt lõi thành công với ba vai. Câu trả lời cuối khớp nội dung lưu và nguồn dẫn. Migration và db:check được kiểm chứng trên đích triển khai. Kết quả tải và đánh giá truy hồi được ghi riêng, với thời điểm, cấu hình và dữ liệu đo.')
h('Tài liệu và mã nguồn đối chiếu',2)
for t in ['docs/PHAM-VI-CNTT.md và docs/phan-quyen.md — phạm vi và quyền hiện hành.','docs/use-cases/README.md và docs/use-cases/sequences/ — danh mục use case và sequence riêng từng chức năng.','server/prisma/schema.prisma và server/prisma/migrations/ — cấu trúc và ràng buộc dữ liệu.','server/src/modules/chat/, retrieval/, documents/, admin/ — nghiệp vụ và giao tiếp API.','web/src/features/ và server/src/netlab/ — giao diện và thực hành mạng.','docker-compose.yml, server/package.json, web/package.json — môi trường và công nghệ.']:p(t)
p('HẾT TÀI LIỆU')
doc.core_properties.title='Phân tích và thiết kế hệ thống Sổ Tay Sinh Viên CNTT'
doc.core_properties.subject='Học phần Lập trình mạng máy tính'
doc.core_properties.author=''
target=ROOT/'docs/Phan-tich-thiet-ke-So-Tay-Sinh-Vien-CNTT.docx'
doc.save(target)
print(target)
