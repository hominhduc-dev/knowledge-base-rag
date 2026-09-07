# -*- coding: utf-8 -*-
"""Khung tài liệu Word cho báo cáo Đồ án Phân tích thiết kế hệ thống.

Định dạng theo mẫu báo cáo môn học: Times New Roman 13, A4, lề trên/dưới 2 cm,
lề trái 3 cm, lề phải 1,5 cm. Hình và bảng đánh số tự động để không lệch khi
chèn thêm nội dung.
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
IMG = Path(__file__).resolve().parent / 'pttkht'

BODY = 'Times New Roman'
MONO = 'Consolas'


class Report:
    def __init__(self):
        doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(3)
        sec.right_margin = Cm(1.5)
        for st in doc.styles:
            for border in list(st.element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
        for name in ['Normal', 'Body Text', 'List Bullet', 'List Number', 'Caption']:
            s = doc.styles[name]
            s.font.name = BODY
            s.font.size = Pt(13)
            s.font.color.rgb = RGBColor(0, 0, 0)
            s.paragraph_format.space_after = Pt(6)
            s.paragraph_format.line_spacing = 1.35
        doc.styles['Caption'].font.size = Pt(12)
        doc.styles['Caption'].font.italic = True
        for name, size in [('Heading 1', 16), ('Heading 2', 14),
                           ('Heading 3', 13), ('Heading 4', 13)]:
            s = doc.styles[name]
            s.font.name = BODY
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.italic = (name == 'Heading 4')
            s.font.color.rgb = RGBColor(0, 0, 0)
            s.paragraph_format.space_before = Pt(12)
            s.paragraph_format.space_after = Pt(6)
            s.paragraph_format.keep_with_next = True
        self.doc = doc
        self.fig_no = 0
        self.tab_no = 0
        self._break = False
        self._footer()

    # -- hạ tầng ------------------------------------------------------------
    def _footer(self):
        f = self.doc.sections[0].footer.paragraphs[0]
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = f.add_run()
        r.font.name = BODY
        r.font.size = Pt(11)
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), 'PAGE')
        f._p.append(fld)

    def page(self):
        self._break = True

    def _apply_break(self, par):
        if self._break:
            par.paragraph_format.page_break_before = True
            self._break = False
        return par

    # -- khối văn bản -------------------------------------------------------
    def p(self, text='', style=None, align=None, bold=False, italic=False,
          size=None, space_after=None):
        par = self.doc.add_paragraph(style=style)
        run = par.add_run(text)
        run.font.name = BODY
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if align is not None:
            par.alignment = align
        else:
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if space_after is not None:
            par.paragraph_format.space_after = Pt(space_after)
        return self._apply_break(par)

    def h(self, text, level=1):
        par = self.doc.add_heading(text, level)
        for r in par.runs:
            r.font.name = BODY
        return self._apply_break(par)

    def bullets(self, items, dash='–'):
        for t in items:
            par = self.doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.45)
            par.paragraph_format.space_after = Pt(3)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = par.add_run(dash + '  ' + t)
            r.font.name = BODY
            self._apply_break(par)

    def steps(self, items, start=1):
        for i, t in enumerate(items, start):
            par = self.doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.45)
            par.paragraph_format.space_after = Pt(3)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = par.add_run(str(i) + '.  ' + t)
            r.font.name = BODY
            self._apply_break(par)

    # -- hình và bảng -------------------------------------------------------
    def fig(self, name, caption, width=15.5):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.space_after = Pt(2)
        self._apply_break(par)
        par.add_run().add_picture(str(IMG / (name + '.png')), width=Cm(width))
        self.fig_no += 1
        cap = self.doc.add_paragraph(style='Caption')
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run('Hình %d. %s' % (self.fig_no, caption))
        r.font.name = BODY
        r.italic = True
        return self.fig_no

    def table(self, headers, rows, widths, caption=None, size=12,
              first_bold=False, header=True):
        if caption:
            self.tab_no += 1
            cap = self.doc.add_paragraph(style='Caption')
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run('Bảng %d. %s' % (self.tab_no, caption))
            r.font.name = BODY
            r.italic = True
            self._apply_break(cap)
        t = self.doc.add_table(rows=(1 if header else 0), cols=len(headers))
        t.style = 'Table Grid'
        t.autofit = False
        if header:
            for c, v in zip(t.rows[0].cells, headers):
                c.text = str(v)
        for row in rows:
            for c, v in zip(t.add_row().cells, row):
                c.text = str(v)
        for ri, row in enumerate(t.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:cantSplit'))
            if ri == 0:
                trPr.append(OxmlElement('w:tblHeader'))
            for ci, c in enumerate(row.cells):
                c.width = Cm(widths[ci])
                for par in c.paragraphs:
                    par.paragraph_format.space_after = Pt(2)
                    par.paragraph_format.space_before = Pt(2)
                    par.paragraph_format.line_spacing = 1.1
                    for r in par.runs:
                        r.font.name = BODY
                        r.font.size = Pt(size)
                        r.bold = (header and ri == 0) or (first_bold and ci == 0)
                if header and ri == 0:
                    sh = OxmlElement('w:shd')
                    sh.set(qn('w:fill'), 'E8EFF4')
                    c._tc.get_or_add_tcPr().append(sh)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        sp.paragraph_format.line_spacing = 1
        return t

    # -- bảng đặc tả ca sử dụng (theo mẫu báo cáo tham chiếu) ---------------
    def uc_spec(self, code, name, actors, pre, post, desc, main, alt, caption=None):
        self.tab_no += 1
        cap = self.doc.add_paragraph(style='Caption')
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run('Bảng %d. Đặc tả ca sử dụng %s — %s' % (self.tab_no, code, name))
        r.font.name = BODY
        r.italic = True
        self._apply_break(cap)
        t = self.doc.add_table(rows=0, cols=1)
        t.style = 'Table Grid'
        t.autofit = False

        def cell(pairs, header=False, numbered=False):
            row = t.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            c = row.cells[0]
            c.width = Cm(15.5)
            c.text = ''
            first = True
            for label, value in pairs:
                par = c.paragraphs[0] if first else c.add_paragraph()
                first = False
                par.paragraph_format.space_after = Pt(2)
                par.paragraph_format.line_spacing = 1.15
                par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                if label:
                    rr = par.add_run(label + ': ')
                    rr.bold = True
                    rr.font.name = BODY
                    rr.font.size = Pt(12)
                rr = par.add_run(value)
                rr.font.name = BODY
                rr.font.size = Pt(12)
                rr.bold = header
            if header:
                sh = OxmlElement('w:shd')
                sh.set(qn('w:fill'), 'E8EFF4')
                c._tc.get_or_add_tcPr().append(sh)

        cell([('Ca sử dụng', code + ' — ' + name),
              ('Các tác nhân', actors),
              ('Điều kiện trước', pre),
              ('Điều kiện sau', post),
              ('Mô tả', desc)])
        cell([('', 'Luồng sự kiện chính')], header=True)
        cell([('', '%d.  %s' % (i, s)) for i, s in enumerate(main, 1)])
        cell([('', 'Luồng sự kiện phụ')], header=True)
        cell([('', s) for s in alt])
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        sp.paragraph_format.line_spacing = 1
        return t

    # -- mã nguồn -----------------------------------------------------------
    def code(self, text, size=10):
        for line in text.rstrip().split('\n'):
            par = self.doc.add_paragraph()
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.line_spacing = 1.05
            par.paragraph_format.left_indent = Cm(0.4)
            r = par.add_run(line if line.strip() else ' ')
            r.font.name = MONO
            r.font.size = Pt(size)
            rPr = r._element.get_or_add_rPr()
            rf = rPr.find(qn('w:rFonts'))
            rf.set(qn('w:eastAsia'), MONO)
            self._apply_break(par)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)
        sp.paragraph_format.line_spacing = 1

    def save(self, path, title, subject):
        self.doc.core_properties.title = title
        self.doc.core_properties.subject = subject
        self.doc.core_properties.author = ''
        self.doc.save(str(path))
        return path


CENTER = WD_ALIGN_PARAGRAPH.CENTER
LEFT = WD_ALIGN_PARAGRAPH.LEFT
JUSTIFY = WD_ALIGN_PARAGRAPH.JUSTIFY
